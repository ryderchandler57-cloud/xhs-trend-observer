from __future__ import annotations

import argparse
import csv
import json
import sqlite3
from pathlib import Path
from typing import Any


def load_json(path: str) -> dict[str, Any]:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def write_json(path: str, data: Any) -> None:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def write_csv(path: str, notes: list[dict[str, Any]]) -> None:
    fields = [
        "keyword",
        "note_id",
        "title",
        "content",
        "author",
        "publish_time",
        "like_count",
        "collect_count",
        "comment_count",
        "images",
        "video",
        "url",
        "extraction_method",
    ]
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for note in notes:
            writer.writerow({field: as_text(note.get(field)) for field in fields})


def write_xlsx(path: str, notes: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    fields = [
        "keyword",
        "note_id",
        "title",
        "content",
        "author",
        "publish_time",
        "like_count",
        "collect_count",
        "comment_count",
        "images",
        "video",
        "url",
        "extraction_method",
    ]
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "notes"
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
    for note in notes:
        ws.append([as_text(note.get(field)) for field in fields])
    widths = {
        "A": 18,
        "B": 24,
        "C": 36,
        "D": 60,
        "E": 18,
        "F": 16,
        "J": 30,
        "K": 30,
        "L": 48,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    comments = wb.create_sheet("valuable_comments")
    comment_fields = ["note_id", "title", "comment", "value_score", "reason", "tags", "url"]
    comments.append(comment_fields)
    for cell in comments[1]:
        cell.font = Font(bold=True)
    for note in notes:
        for comment in note.get("valuable_comments", []) or []:
            comments.append(
                [
                    note.get("note_id", ""),
                    note.get("title", ""),
                    comment.get("comment", ""),
                    comment.get("value_score", 0),
                    comment.get("reason", ""),
                    as_text(comment.get("tags", [])),
                    note.get("url", ""),
                ]
            )
    for column, width in {"A": 24, "B": 36, "C": 60, "E": 48, "F": 24, "G": 48}.items():
        comments.column_dimensions[column].width = width
    for row in comments.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(out)


def write_sqlite(path: str, notes: list[dict[str, Any]]) -> None:
    db = Path(path)
    db.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(db)
    try:
        conn.execute("DROP TABLE IF EXISTS notes")
        conn.execute("DROP TABLE IF EXISTS comments")
        conn.execute(
            """
            CREATE TABLE notes (
              note_id TEXT PRIMARY KEY,
              keyword TEXT,
              title TEXT,
              content TEXT,
              author TEXT,
              publish_time TEXT,
              like_count INTEGER,
              collect_count INTEGER,
              comment_count INTEGER,
              images TEXT,
              video TEXT,
              url TEXT,
              extraction_method TEXT
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE comments (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              note_id TEXT,
              comment TEXT,
              value_score INTEGER,
              reason TEXT,
              tags TEXT
            )
            """
        )
        for note in notes:
            conn.execute(
                """
                INSERT OR REPLACE INTO notes VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    note.get("note_id", ""),
                    note.get("keyword", ""),
                    note.get("title", ""),
                    note.get("content", ""),
                    note.get("author", ""),
                    note.get("publish_time", ""),
                    note.get("like_count"),
                    note.get("collect_count"),
                    note.get("comment_count"),
                    as_text(note.get("images", [])),
                    note.get("video", ""),
                    note.get("url", ""),
                    note.get("extraction_method", ""),
                ),
            )
            for comment in note.get("valuable_comments", []) or []:
                conn.execute(
                    "INSERT INTO comments(note_id, comment, value_score, reason, tags) VALUES (?, ?, ?, ?, ?)",
                    (
                        note.get("note_id", ""),
                        comment.get("comment", ""),
                        comment.get("value_score", 0),
                        comment.get("reason", ""),
                        as_text(comment.get("tags", [])),
                    ),
                )
        conn.commit()
    finally:
        conn.close()


def markdown(report: dict[str, Any]) -> str:
    meta = report.get("meta", {})
    summary = report.get("ai_summary", {}) or {}
    lines = [
        f"# {meta.get('title', '趋势与认知观察报告')}",
        "",
        f"- 搜索关键词：{meta.get('query', '')}",
        f"- 主题：{meta.get('topic', '')}",
        f"- 关注方向：{' / '.join(meta.get('focus_topics', []) or [])}",
        f"- 生成时间：{meta.get('generated_at', '')}",
        "",
        "## AI总结",
    ]
    for title, key in [
        ("高频观点", "high_frequency_views"),
        ("用户真实痛点", "user_pain_points"),
        ("避坑点", "avoid_pitfalls"),
        ("可执行建议", "actionable_suggestions"),
        ("内容结论", "content_conclusions"),
    ]:
        lines.extend(["", f"### {title}"])
        values = [item for item in summary.get(key, []) if item]
        lines.extend([f"- {item}" for item in values] or ["- 暂无足够证据"])
    lines.extend(["", "### KOL/博主分析"])
    for item in summary.get("kol_analysis", [])[:10]:
        lines.append(f"- {item.get('author', '')}：{item.get('note_count', 0)} 篇，点赞 {item.get('total_likes', 0)}")
    lines.extend(["", "### 评论情绪"])
    sentiment = summary.get("comment_sentiment", {}).get("counts", {})
    lines.append(f"- 正向：{sentiment.get('positive', 0)}")
    lines.append(f"- 负向：{sentiment.get('negative', 0)}")
    lines.append(f"- 中性：{sentiment.get('neutral', 0)}")
    lines.extend(["", "### 爆款结构拆解"])
    for item in summary.get("viral_note_teardown", [])[:10]:
        lines.append(f"- {item.get('title', '')}：{' / '.join(item.get('hooks', []))}")
    lines.extend(["", "### 竞品/渠道对比"])
    for item in summary.get("competitor_comparison", [])[:10]:
        lines.append(f"- {' / '.join(item.get('entities', []))}：{item.get('note_title', '')}")
    lines.extend(["", "## 高价值评论榜"])
    for item in report.get("high_value_comments", [])[:30]:
        lines.append(f"- [{item.get('value_score', 0)}分] {item.get('comment', '')}")
        if item.get("reason"):
            lines.append(f"  - 理由：{item['reason']}")
        if item.get("note_url"):
            lines.append(f"  - 来源：{item['note_url']}")
    lines.extend(["", "## 原文链接"])
    for note in report.get("structured_notes", [])[:80]:
        if note.get("url"):
            lines.append(f"- {note.get('title') or note.get('note_id')}: {note['url']}")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Export structured XHS report artifacts.")
    parser.add_argument("--report", required=True)
    parser.add_argument("--notes-json", required=True)
    parser.add_argument("--csv", required=True)
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sqlite", required=True)
    parser.add_argument("--markdown", required=True)
    args = parser.parse_args()

    report = load_json(args.report)
    notes = report.get("structured_notes", []) or []
    write_json(args.notes_json, notes)
    write_csv(args.csv, notes)
    write_xlsx(args.xlsx, notes)
    write_sqlite(args.sqlite, notes)
    Path(args.markdown).parent.mkdir(parents=True, exist_ok=True)
    Path(args.markdown).write_text(markdown(report), encoding="utf-8")
    print(f"structured notes JSON written: {args.notes_json}")
    print(f"CSV written: {args.csv}")
    print(f"Excel written: {args.xlsx}")
    print(f"SQLite written: {args.sqlite}")
    print(f"Markdown written: {args.markdown}")


if __name__ == "__main__":
    main()
